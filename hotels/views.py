import csv
import io

from django.contrib import messages
from django.contrib.auth.decorators import permission_required
from django.db import transaction
from django.db.models.functions import Lower
from django.shortcuts import render, redirect, get_object_or_404
from django.views.generic import ListView
from .models import Hotel
from .forms import HotelImportForm


class HotelListView(ListView):
    """
    Список отелей с поиском по городу, сортировкой и постраничным выводом.

    До аудита это была простая функция без пагинации и без сортировки —
    единственное место в проекте, где каталог ещё выводился целиком одним
    списком (аудит, дефект №8 «нет пагинации в списках отелей и номеров» —
    для rooms это уже было исправлено на Занятии 1, для hotels — нет).
    Переведено на тот же подход, что и RoomListView в rooms/views.py, для
    единообразия: CBV + paginate_by + {% querystring %} в шаблоне.
    """
    model = Hotel
    template_name = 'hotels/list.html'
    context_object_name = 'hotels'
    paginate_by = 6

    # Ключи ?sort= из формы → реальные выражения order_by().
    # '-rating' (сначала с лучшим рейтингом) — сортировка по умолчанию.
    SORTS = {
        'price': 'min_price',
        '-price': '-min_price',
        'rating': '-rating',
    }

    def get_queryset(self):
        # with_stats() даёт min_price (честная цена «от», посчитанная по
        # реальным номерам отеля — аудит, дефект №14) и rooms_count.
        hotels = Hotel.objects.active().with_stats()

        city = self.request.GET.get('city', '')
        if city:
            # Стандартный icontains в SQLite регистронезависим только для
            # латиницы, поэтому для кириллицы приводим и поле, и запрос
            # к нижнему регистру через Lower() (Unicode-версию регистрирует
            # core/apps.py::CoreConfig.ready()).
            hotels = hotels.annotate(city_lower=Lower('city')).filter(
                city_lower__contains=city.lower()
            )

        sort_by = self.request.GET.get('sort', '')
        return hotels.order_by(self.SORTS.get(sort_by, '-rating'))

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['city'] = self.request.GET.get('city', '')
        context['sort'] = self.request.GET.get('sort', '')
        return context


def hotel_detail(request, pk):
    """
    Страница одного отеля со списком его номеров.

    get_object_or_404 вернёт отель по первичному ключу или отдаст страницу
    404, если такого отеля нет. Номера берём через обратную связь
    hotel.rooms (related_name из модели Room).
    """
    # with_stats() — тот же агрегат min_price, что и в списке отелей
    # (HotelListView), чтобы цена «от» совпадала на обеих страницах.
    hotel = get_object_or_404(Hotel.objects.with_stats(), pk=pk, is_active=True)
    # on_sale()/with_today_status() — см. rooms/models.py::RoomQuerySet;
    # без with_today_status() карточка номера не могла бы показать бейдж
    # «Занят сейчас», полагаясь на статичное is_available (дефект №4).
    rooms = hotel.rooms.on_sale().with_today_status()
    return render(request, 'hotels/detail.html', {
        'hotel': hotel,
        'rooms': rooms,
    })


# Обязательные колонки в файле импорта, в этом порядке.
IMPORT_COLUMNS = ['name', 'city', 'address', 'price', 'rating']


def _rows_from_csv(uploaded_file):
    """Строки из CSV как список словарей {колонка: значение}."""
    decoded = io.TextIOWrapper(uploaded_file.file, encoding='utf-8-sig')
    return list(csv.DictReader(decoded))


def _rows_from_xlsx(uploaded_file):
    """
    Строки из Excel как список словарей — тем же форматом, что и CSV,
    чтобы дальше по коду формат файла уже был не важен: одна и та же
    функция создания отелей работает что с CSV-, что с Excel-строками.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    sheet = workbook.active
    rows_iter = sheet.iter_rows(values_only=True)

    header = [str(cell).strip() if cell is not None else '' for cell in next(rows_iter)]
    rows = []
    for values in rows_iter:
        if all(v is None for v in values):  # пропускаем полностью пустые строки
            continue
        row = dict(zip(header, values))
        rows.append({k: ('' if v is None else str(v)) for k, v in row.items()})
    return rows


@permission_required('hotels.add_hotel', raise_exception=True)
def import_hotels(request):
    """
    Импорт отелей из CSV или Excel (.xlsx). Доступно только персоналу с
    правом hotels.add_hotel (аудит: «Разграничение доступа» из плана
    занятия 2) — обычный посетитель сайта не должен иметь возможность
    массово добавлять отели через загрузку файла.

    raise_exception=True — вместо редиректа на страницу входа сразу
    отдаём 403 Forbidden: для страницы импорта это понятнее, чем форма
    входа, которую всё равно не поможет пройти обычному пользователю.

    Формат файла определяется по расширению, а не угадывается по
    содержимому — дальше оба формата превращаются в одинаковый список
    словарей (_rows_from_csv / _rows_from_xlsx), и вся остальная логика
    (валидация, создание отелей, транзакция) общая для обоих форматов.

    Вся операция обёрнута в transaction.atomic(): либо создаются ВСЕ
    строки файла, либо (при ошибке в любой из них) ни одной — так импорт
    не оставляет базу в наполовину загруженном состоянии.
    """
    if request.method == 'POST':
        form = HotelImportForm(request.POST, request.FILES)
        if form.is_valid():
            uploaded_file = form.cleaned_data['import_file']
            if uploaded_file.name.lower().endswith('.xlsx'):
                rows = _rows_from_xlsx(uploaded_file)
            else:
                rows = _rows_from_csv(uploaded_file)

            created_rows = []
            errors = []
            try:
                with transaction.atomic():
                    for line_number, row in enumerate(rows, start=2):  # строка 1 — заголовок
                        missing = [c for c in IMPORT_COLUMNS if not row.get(c)]
                        if missing:
                            errors.append(f'Строка {line_number}: не хватает полей {", ".join(missing)}')
                            continue
                        try:
                            hotel = Hotel(
                                name=row['name'].strip(),
                                city=row['city'].strip(),
                                address=row['address'].strip(),
                                price=row['price'],
                                rating=row['rating'],
                            )
                            hotel.full_clean()  # прогоняет те же валидаторы, что и форма/админка
                        except Exception as exc:
                            errors.append(f'Строка {line_number}: {exc}')
                            continue
                        hotel.save()
                        created_rows.append(hotel.name)

                    if errors:
                        # Откатываем всю транзакцию — либо всё, либо ничего.
                        raise ValueError('import aborted: validation errors')
            except ValueError:
                for err in errors:
                    messages.error(request, err)
                messages.error(request, 'Импорт отменён: ни один отель не создан из-за ошибок выше.')
            else:
                messages.success(request, f'Импортировано отелей: {len(created_rows)}.')
                return redirect('hotels:hotel_list')
    else:
        form = HotelImportForm()

    return render(request, 'hotels/import.html', {'form': form})
