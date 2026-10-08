# ────────────────────────────────────────────────────────────────────────────
# hotels/exports.py — выгрузка каталога отелей в нескольких форматах:
# CSV, Excel (.xlsx), PDF и Word (.docx).
#
# Занятие 2, «Экспорт и импорт данных»: один и тот же набор данных отдаём в
# четырёх форматах — это наглядно показывает разницу между «форматом для
# программ» (CSV, Excel — их читают другие системы и таблицы) и «форматом
# для людей» (PDF, DOCX — открывают и читают глазами, распечатывают).
#
# Реализация каждого формата вынесена в отдельную маленькую функцию, а не
# размазана по одной большой — так легче понять, какая библиотека за что
# отвечает: csv (стандартная библиотека), openpyxl (Excel), reportlab (PDF),
# python-docx (Word).
# ────────────────────────────────────────────────────────────────────────────
import csv
import io

from django.http import HttpResponse

from .models import Hotel

# Заголовки колонок — общие для всех форматов, чтобы выгрузка выглядела
# одинаково независимо от того, что выбрал пользователь.
COLUMNS = ['ID', 'Название', 'Город', 'Адрес', 'Цена за ночь', 'Рейтинг']


def _hotel_rows():
    """Общий источник строк для всех форматов экспорта — один запрос к БД,
    без повторения фильтра `is_active=True` в каждой функции."""
    for hotel in Hotel.objects.filter(is_active=True).order_by('name'):
        yield [hotel.pk, hotel.name, hotel.city, hotel.address, hotel.price, hotel.rating]


def export_hotels_csv(request):
    """Экспорт в CSV — для импорта в другую систему или в тот же HotelHub."""
    response = HttpResponse(content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = 'attachment; filename="hotels.csv"'

    # \ufeff — BOM: без него Excel на Windows иногда показывает кириллицу
    # в CSV как набор «кракозябр» вместо букв.
    response.write('\ufeff')

    writer = csv.writer(response, delimiter=';')
    writer.writerow(COLUMNS)
    for row in _hotel_rows():
        writer.writerow(row)
    return response


def export_hotels_xlsx(request):
    """
    Экспорт в Excel (.xlsx) через openpyxl.

    В отличие от CSV, настоящий .xlsx хранит типы данных (число, а не
    текст «45000»), поддерживает ширину колонок и жирный заголовок —
    поэтому его удобнее открывать и сразу читать глазами в Excel/LibreOffice.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = 'Отели'

    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in _hotel_rows():
        ws.append(row)

    # Автоширина колонок «на глаз» — просто чуть шире самого длинного значения.
    for col_cells in ws.columns:
        length = max(len(str(c.value)) for c in col_cells if c.value is not None)
        ws.column_dimensions[col_cells[0].column_letter].width = length + 4

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = 'attachment; filename="hotels.xlsx"'
    return response


def export_hotels_pdf(request):
    """
    Экспорт в PDF через reportlab — «печатная» версия каталога.

    PDF — формат для чтения и печати, а не для повторной загрузки данных
    обратно в систему, поэтому здесь нет parsing-логики, только вывод.
    """
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                             leftMargin=1.5 * cm, rightMargin=1.5 * cm,
                             topMargin=1.5 * cm, bottomMargin=1.5 * cm)

    styles = getSampleStyleSheet()
    elements = [Paragraph('HotelHub — каталог отелей', styles['Title']), Spacer(1, 12)]

    data = [COLUMNS] + [[str(v) for v in row] for row in _hotel_rows()]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#161B2E')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F4F6FC')]),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
    ]))
    elements.append(table)
    doc.build(elements)
    buffer.seek(0)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="hotels.pdf"'
    return response


def export_hotels_docx(request):
    """
    Экспорт в Word (.docx) через python-docx — редактируемый документ,
    удобный, например, персоналу для распечатки или правки офлайн.
    """
    from docx import Document

    document = Document()
    document.add_heading('HotelHub — каталог отелей', level=1)

    table = document.add_table(rows=1, cols=len(COLUMNS))
    table.style = 'Light Grid Accent 1'
    header_cells = table.rows[0].cells
    for i, col_name in enumerate(COLUMNS):
        header_cells[i].text = col_name

    for row in _hotel_rows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)

    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    response['Content-Disposition'] = 'attachment; filename="hotels.docx"'
    return response
